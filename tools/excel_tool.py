"""generate_testcase_excel 工具:把 Markdown 用例表转换为带样式的 xlsx 并落盘。

由 FileUploadMiddleware 注册到 agent(见 middleware/file_upload.py),
这里只做纯粹的 表格解析 -> Excel 生成。

双写:除落真实磁盘 outputs/ 外,同时以 Command 把 Markdown 终稿(utf-8)和
xlsx(base64)写入 agent 虚拟文件系统的 /outputs/ 路径,前端文件面板从
state.files 读取即可看到并下载(与 generate_xmind 的行为一致)。
"""

from __future__ import annotations

import base64
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Annotated

from langchain_core.messages import ToolMessage
from langchain_core.tools import InjectedToolCallId, tool
from langgraph.types import Command
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Excel 输出目录(agent 项目根下的 outputs/)
OUTPUT_DIR = Path(__file__).resolve().parent.parent / "outputs"

EXPECTED_HEADERS = ["用例编号", "用例名称", "前置条件", "测试步骤", "预期结果", "优先级"]

_PRIORITY_FILL = {
    "P0": PatternFill("solid", fgColor="F8CBAD"),  # 橙红
    "P1": PatternFill("solid", fgColor="FFE699"),  # 黄
    "P2": PatternFill("solid", fgColor="C6E0B4"),  # 绿
    "P3": PatternFill("solid", fgColor="D9D9D9"),  # 灰
}

_THIN_BORDER = Border(
    *(Side(style="thin", color="BFBFBF"),) * 4,
)

_COLUMN_WIDTHS = [16, 28, 24, 48, 48, 10]


class MarkdownTableError(ValueError):
    """Markdown 用例表不符合约定格式。"""


def _safe_name(doc_name: str) -> str:
    """文件名安全化:替换非法字符,空名兜底为 testcases。"""
    return re.sub(r'[\\/:*?"<>|]+', "_", doc_name).strip() or "testcases"


def parse_markdown_table(markdown_table: str) -> list[list[str]]:
    """解析 Markdown 表格为二维数组(含表头行)。

    约定:第一行是表头,第二行是 |---|---| 分隔行,之后每行一条用例。
    单元格内不允许出现未转义的 `|`(由提示词约束)。
    行首/行尾的 `|` 可有可无(模型时常省略),按竖线数识别表格行。
    """
    lines = [ln.strip() for ln in markdown_table.strip().splitlines() if ln.strip()]
    # 定位表头(列名与顺序完全一致),只吃紧随其后的数据行;
    # 遇到非表格行即停止,避免把文末附录的矩阵/判定表混进来
    header_idx = None
    for i, ln in enumerate(lines):
        if ln.count("|") >= len(EXPECTED_HEADERS) - 1:
            cells = [c.strip() for c in ln.strip("|").split("|")]
            if cells == EXPECTED_HEADERS:
                header_idx = i
                break
    rows: list[list[str]] = []
    if header_idx is not None:
        rows.append(EXPECTED_HEADERS)
        for ln in lines[header_idx + 1 :]:
            if ln.count("|") < len(EXPECTED_HEADERS) - 1:
                break  # 表格结束
            cells = [c.strip() for c in ln.strip("|").split("|")]
            # 跳过分隔行 |---|---|
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue
            rows.append(cells)

    if not rows:
        raise MarkdownTableError("未在输入中找到 Markdown 表格(行需以 | 分隔六列)")
    header = rows[0]
    if header[: len(EXPECTED_HEADERS)] != EXPECTED_HEADERS:
        raise MarkdownTableError(
            f"表头不符合约定,期望:{' | '.join(EXPECTED_HEADERS)},实际:{' | '.join(header)}"
        )
    data = [r for r in rows[1:] if any(c for c in r)]
    if not data:
        raise MarkdownTableError("表格中没有用例数据行")
    return [header, *data]


def write_excel(rows: list[list[str]], doc_name: str) -> Path:
    """把用例二维数组写入带样式的 xlsx,返回文件路径。"""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUTPUT_DIR / f"{_safe_name(doc_name)}-testcases.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "测试用例"

    for r_idx, row in enumerate(rows, start=1):
        for c_idx in range(1, len(EXPECTED_HEADERS) + 1):
            value = row[c_idx - 1] if c_idx - 1 < len(row) else ""
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(
                vertical="center",
                horizontal="center" if r_idx == 1 or c_idx in (1, 6) else "left",
                wrap_text=True,
            )
            if r_idx == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="4472C4")

    # 优先级列配色(数据行)
    priority_col = EXPECTED_HEADERS.index("优先级") + 1
    for r_idx in range(2, len(rows) + 1):
        cell = ws.cell(row=r_idx, column=priority_col)
        fill = _PRIORITY_FILL.get(str(cell.value).strip().upper())
        if fill:
            cell.fill = fill

    for i, width in enumerate(_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"

    wb.save(path)
    return path


def write_markdown(markdown_table: str, doc_name: str) -> Path:
    """把用例 Markdown 表格原样落盘为终稿 md,返回文件路径。"""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    safe_name = _safe_name(doc_name)
    path = OUTPUT_DIR / f"{safe_name}-testcases.md"
    path.write_text(
        f"# {safe_name} 测试用例\n\n{markdown_table.strip()}\n", encoding="utf-8"
    )
    return path


def _to_file_data(content: str, encoding: str = "utf-8") -> dict:
    """构造 deepagents 虚拟文件系统的 FileData 对象(与前端 types.ts 结构一致)。"""
    now = datetime.now(timezone.utc).isoformat()
    return {"content": content, "encoding": encoding,
            "created_at": now, "modified_at": now}


@tool
def generate_testcase_excel(
    markdown_table: str,
    doc_name: str,
    tool_call_id: Annotated[str, InjectedToolCallId],
) -> Command | str:
    """把完整的测试用例 Markdown 表格转换为 Excel(.xlsx)文件并保存到本机,
    同时把表格原样保存为 Markdown 终稿(.md)。

    Args:
        markdown_table: 完整用例表,列必须为「用例编号 | 用例名称 | 前置条件 | 测试步骤 | 预期结果 | 优先级」,
            包含表头行与分隔行,单元格内不得包含 `|` 或换行。
        doc_name: 需求文档名(不含扩展名),用于命名输出文件。

    Returns:
        保存成功时返回 xlsx 与 md 的绝对路径;失败时返回以 ERROR: 开头的错误说明。
    """
    try:
        rows = parse_markdown_table(markdown_table)
        xlsx_path = write_excel(rows, doc_name)
        md_path = write_markdown(markdown_table, doc_name)
    except MarkdownTableError as e:
        return f"ERROR: Markdown 用例表格式错误:{e}。请修正表格后重新调用本工具。"
    except Exception as e:  # noqa: BLE001 - 工具不能把异常抛给模型循环
        return f"ERROR: 生成 Excel 失败:{type(e).__name__}: {e}"

    summary = f"Excel 用例表: {xlsx_path}\nMarkdown 终稿: {md_path}"
    # 双写:除落真实磁盘,同时写入 agent 虚拟文件系统的 /outputs/ 路径,
    # 前端文件面板从 state.files 读取展示;xlsx 为二进制,按 base64 编码入账
    # (与文件上传注入的 FileData 结构一致)。files 状态走合并 reducer,
    # 只追加这两个 key,不影响 /uploads/ 下的文件。
    return Command(
        update={
            "files": {
                f"/outputs/{md_path.name}": _to_file_data(
                    md_path.read_text(encoding="utf-8")
                ),
                f"/outputs/{xlsx_path.name}": _to_file_data(
                    base64.b64encode(xlsx_path.read_bytes()).decode("ascii"),
                    encoding="base64",
                ),
            },
            "messages": [ToolMessage(content=summary, tool_call_id=tool_call_id)],
        }
    )
